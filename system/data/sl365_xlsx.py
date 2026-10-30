import sys
sys.path.insert(0, "/home/claude")
from sl365_build import rows, PROD, BOARD_NAMES, NEW_BOARDS, BOARDS
from blog_seo import out as SEO
SEOD = {o["day"]: o for o in SEO}
for r in rows:  # pin copy now follows the post's primary keyword
    o = SEOD[r["day"]]
    sec = o["secondary"].split("; ")
    r.update(phrase1=o["primary"], phrase2=sec[0], phrase3=sec[1] if len(sec) > 1 else "",
             pin_title=o["pin_title"], pin_desc=o["pin_desc"], alt=o["alt"], post_title=o["h1"], slug=o["slug"])
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

F = "Arial"
NAVY = "14213D"
HF = Font(name=F, bold=True, color="FFFFFF", size=10)
H = PatternFill("solid", start_color=NAVY)
BF = Font(name=F, size=10)
LF = Font(name=F, size=10, color="0563C1", underline="single")
wb = Workbook()

def table(ws, headers, data, widths, name, link_cols=(), wrap_cols=()):
    ws.append(headers)
    for r in data:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = HF, H
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
            if c.column in link_cols and c.value:
                c.hyperlink = c.value
                c.font = LF
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    t = Table(displayName=name, ref=f"A1:{get_column_letter(len(headers))}{len(data)+1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)

# ---- Start Here ----
ws = wb.active
ws.title = "Start Here"
lines = [
 ("Sideline Ready - 365-Day Product & Pin Plan", 14, True),
 ("One anchor product per day, Dec 1 2026 - Nov 30 2027. Each day = 1 blog post (story or how-to) + its pins + a 5-7 item Amazon bundle of $100 or more.", 10, False),
 ("", 10, False),
 ("How each day works", 11, True),
 ("1. Mom searches Pinterest using the exact phrase in 'Pinterest search phrase'.", 10, False),
 ("2. She sees the pin (title, description and alt text are drafted in this sheet) and clicks through to the blog post.", 10, False),
 ("3. The post tells the story / solves the problem, then ends with a call to action: the full bundle with an 'add all to cart' button.", 10, False),
 ("4. Anything she adds to her Amazon cart within 24 hours of clicking counts for you, and stays credited for 89 days if she checks out later.", 10, False),
 ("", 10, False),
 ("The customer journey (Blog SEO tab)", 11, True),
 ("Problem: she has a problem right now (cold bleachers, smelly cleats, 95-degree field). Post = how-to / survival guide; the bundle is the fix.", 10, False),
 ("Planning: she's getting ready (packing list, team mom duties, snack ideas). Post = checklist / ideas list; the bundle is the checklist.", 10, False),
 ("Comparing: she's deciding what to buy (pitching machine, wagon, stadium seat). Post = setup or buying guide; the bundle is the setup.", 10, False),
 ("Ready to buy: she's shopping for a gift. Post = gift guide; the bundle is the gift.", 10, False),
 ("Each post: SEO title for Google, H1 that doubles as the pin title, an opening story hook, then the bundle call to action at the end.", 10, False),
 ("", 10, False),
 ("Product rules (checked Oct 30, 2026)", 11, True),
 ("Under $100: 1,000+ bought last month on Amazon and 4.2+ stars. $100 and up: 100+ bought last month and 4.3+ stars. No sponsored listings.", 10, False),
 ("Search phrases are real Pinterest search suggestions pulled Oct 30, 2026 - the words moms actually type.", 10, False),
 ("Each month before I write that month's posts, I re-check prices, stars and sales, swap anything that dropped, and finalize the bundle.", 10, False),
 ("", 10, False),
 ("Drafts vs. final", 11, True),
 ("Pin title, description, alt text, blog title and bundle are DRAFTS. We finalize them together each month when the posts are written.", 10, False),
 ("Cart total uses Oct 30 prices. The 'Add-all-to-cart link' builds the whole bundle on Amazon in one click - test a few before relying on it.", 10, False),
 ("", 10, False),
 ("Before the first post goes live", 11, True),
 ("Create the 3 new Pinterest boards on the 'New Boards' tab.", 10, False),
 ("Add the blog's web address to your Amazon Associates account (Account Settings > Website and mobile app list).", 10, False),
 ("Every post needs the line 'As an Amazon Associate I earn from qualifying purchases' near the top.", 10, False),
 ("Links use your tracking ID screenfreet0f-20. Send a Sideline Ready tracking ID anytime and I'll swap them all.", 10, False),
]
for i, (t, size, bold) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=size, bold=bold, color=NAVY if bold else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 125

# ---- 365 plan ----
ws = wb.create_sheet("365-Day Plan")
hdr = ["Day", "Date", "Month", "Season", "Anchor product", "Price", "Stars", "Ratings", "Bought last month",
       "Pinterest search phrase", "Phrase 2", "Phrase 3", "Pinterest board",
       "Pin title (draft)", "Pin description (draft)", "Alt text (draft)", "Blog post title (draft)", "Post URL slug",
       "Bundle add-ons", "Items in cart", "Cart total", "ASIN", "Anchor affiliate link", "Add-all-to-cart link", "Source"]
data = [[r["day"], r["date"], r["month"], r["season"], r["product"], r["price"], r["rating"], r["ratings"], r["bought"],
         r["phrase1"], r["phrase2"], r["phrase3"], r["board"], r["pin_title"], r["pin_desc"], r["alt"], r["post_title"],
         r["slug"], r["bundle"], r["items"], r["cart_total"], r["asin"], r["link"], r["cart_link"], r["source"]] for r in rows]
table(ws, hdr, data, [6, 11, 10, 9, 42, 10, 7, 9, 11, 32, 30, 30, 36, 50, 70, 60, 60, 40, 80, 8, 11, 13, 50, 60, 16],
      "Plan365", link_cols=(23, 24), wrap_cols=())
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    for c in row:
        c.number_format = '"$"#,##0.00'
for row in ws.iter_rows(min_row=2, min_col=21, max_col=21):
    for c in row:
        c.number_format = '"$"#,##0.00'
ws.freeze_panes = "F2"
ws["U1"].comment = Comment("Anchor + add-ons at Oct 30, 2026 Amazon prices. Re-checked monthly.", "Claude")

# ---- Blog SEO ----
ws = wb.create_sheet("Blog SEO", 2)
data = [[o["day"], o["date"], o["product"], o["stage"], o["intent"], o["ptype"], o["primary"], o["secondary"],
         o["seo_title"], len(o["seo_title"]), o["meta"], len(o["meta"]), o["h1"], "/" + o["slug"], o["hook"], o["cta"]] for o in SEO]
table(ws, ["Day", "Date", "Anchor product", "Journey stage", "Search intent", "Post type", "Primary keyword (exact Pinterest search)",
           "Secondary keywords", "SEO title (Google)", "Chars", "Meta description", "Chars", "H1 / Pin title", "URL slug",
           "Story hook (opening line)", "Bundle call to action"],
      data, [6, 11, 38, 13, 13, 13, 34, 48, 58, 6, 80, 6, 70, 48, 80, 70], "BlogSEO", wrap_cols=())
ws.freeze_panes = "D2"
ws["J1"].comment = Comment("Google shows about 60 characters of a title.", "Claude")
ws["L1"].comment = Comment("Google shows about 155 characters of a meta description.", "Claude")

# ---- Product list (all products) ----
ws = wb.create_sheet("All Products")
day_of = {r["asin"]: r["day"] for r in rows}
data = []
for p in sorted(PROD, key=lambda p: (BOARD_NAMES[p["board"]], -p["price"])):
    data.append([p["name"], p["price"], p["rating"], p["ratings"], p["bought"],
                 {1: "Nov-Feb", 2: "Year-round", 3: "Mar-Aug"}[p["tier"]], BOARD_NAMES[p["board"]],
                 p["kws"][0], "; ".join(p["kws"][1:]), day_of.get(p["asin"], ""), p["asin"],
                 f"https://www.amazon.com/dp/{p['asin']}?tag=screenfreet0f-20", p["src"]])
table(ws, ["Product", "Price", "Stars", "Ratings", "Bought last month", "Best season", "Pinterest board",
           "Pinterest search phrase", "More phrases", "Anchor day #", "ASIN", "Affiliate link", "Source"],
      data, [46, 10, 7, 9, 11, 12, 40, 34, 60, 10, 13, 50, 16], "AllProducts", link_cols=(12,))
for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
    for c in row:
        c.number_format = '"$"#,##0.00'
ws.freeze_panes = "B2"

# ---- New boards ----
ws = wb.create_sheet("New Boards")
table(ws, ["Board name (create in Pinterest)", "Board description"],
      [[BOARD_NAMES[k], d] for k, d in NEW_BOARDS.items()], [48, 120], "NewBoards", wrap_cols=(2,))

# ---- Summary (formulas) ----
ws = wb.create_sheet("Summary")
ws["A1"] = "Summary"; ws["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
stats = [
 ("Days planned", "=COUNTA('365-Day Plan'!A2:A400)"),
 ("Unique anchor products", "=SUMPRODUCT(1/COUNTIF('365-Day Plan'!V2:V366,'365-Day Plan'!V2:V366))"),
 ("Smallest cart ($)", "=MIN('365-Day Plan'!U2:U366)"),
 ("Average cart ($)", "=AVERAGE('365-Day Plan'!U2:U366)"),
 ("Largest cart ($)", "=MAX('365-Day Plan'!U2:U366)"),
 ("Carts of $100 or more", "=COUNTIF('365-Day Plan'!U2:U366,\">=100\")"),
 ("Anchors priced $100+", "=COUNTIF('365-Day Plan'!F2:F366,\">=100\")"),
]
for i, (k, f) in enumerate(stats, 3):
    ws.cell(row=i, column=1, value=k).font = BF
    c = ws.cell(row=i, column=2, value=f); c.font = BF
    if "$" in k:
        c.number_format = '"$"#,##0.00'
r0 = len(stats) + 5
ws.cell(row=r0, column=1, value="Posts per board").font = Font(name=F, bold=True, color=NAVY)
for c, v in zip((1, 2), ("Board", "Posts")):
    x = ws.cell(row=r0 + 1, column=c, value=v); x.font, x.fill = HF, H
for i, name in enumerate(sorted(set(BOARD_NAMES.values())), r0 + 2):
    ws.cell(row=i, column=1, value=name).font = BF
    ws.cell(row=i, column=2, value=f"=COUNTIF('365-Day Plan'!M:M,A{i})").font = BF
end = r0 + 2 + len(set(BOARD_NAMES.values()))
ws.cell(row=end, column=1, value="Total").font = Font(name=F, bold=True)
ws.cell(row=end, column=2, value=f"=SUM(B{r0+2}:B{end-1})").font = Font(name=F, bold=True)
m0 = end + 2
ws.cell(row=m0, column=1, value="Posts per month").font = Font(name=F, bold=True, color=NAVY)
months = []
for r in rows:
    if r["month"] not in months:
        months.append(r["month"])
for i, m in enumerate(months, m0 + 1):
    ws.cell(row=i, column=1, value=m).font = BF
    ws.cell(row=i, column=2, value=f"=COUNTIF('365-Day Plan'!C:C,A{i})").font = BF
ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 14

out = "/home/claude/Sideline_Ready_365_Product_Plan.xlsx"
wb.save(out)
print(out, len(rows))
