"""Sideline Ready big-ticket product list: $100+, 4.0+ stars, 100+ reviews (Amazon scan Oct 30, 2026)."""
import glob, re, sys
sys.path.insert(0, "/home/claude")
from bt_meta import QX, tag
from sl365_build import PROD as PLAN
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TAG = "screenfreet0f-20"
BOARDS = {"train": "Backyard Sports Training Equipment", "gifts": "Sports Mom Gifts & Gift Baskets", "seat": "Stadium Seats & Bleacher Comfort",
          "wagon": "Sports Wagon Ideas & Accessories", "travel": "Travel Ball & Tournament Must Haves", "cold": "Cold Football Game Essentials",
          "bag": "Sports Mom Must Haves & Bag Essentials", "dad": "Sports Dad Essentials & Gifts", "store": "Sports Equipment Storage & Garage Organization",
          "gameday": "Game Day Bags & Stadium Bags", "coach": "Coach Gifts & Team Mom Ideas",
          "fuel": "Teen Athlete Meals, Snacks & Meal Prep"}
in_plan = {p["asin"] for p in PLAN}

def num(b):
    m = re.match(r"([\d.]+)([KM]?)", b or "")
    return float(m.group(1)) * {"K": 1e3, "M": 1e6, "": 1}[m.group(2)] if m else 0

rows, seen = [], set()
for f in sorted(glob.glob("/home/claude/bt_raw/c*.txt")):
    for line in open(f):
        line = line.strip()
        if not line:
            continue
        a, p, r, n, b, q, t = line.split("|", 6)
        if a in seen:
            continue
        seen.add(a)
        p, r, n, q = float(p), float(r), int(float(n)), int(q)
        assert p >= 100 and r >= 4.0 and n >= 100, line
        group, sport, board = tag(QX[q])
        OVR = {"batting helmet": ("Player gear (gift-worthy)", "gifts"), "football shoulder pads youth": ("Player gear (gift-worthy)", "gifts"),
               "pickleball paddle": ("Player gear (gift-worthy)", "gifts"), "ipad": ("Tech gifts", "gifts"),
               "scooter electric teen": ("Gifts for kids & teens", "gifts"), "heated blanket battery": ("Cold-weather gear", "cold"),
               "heated gloves": ("Cold-weather gear", "cold"), "basketball arcade game": ("Game room & arcade", "gifts"),
               "blender": ("Athlete meals (blenders)", "fuel"), "golf cart": ("Golf (sports dad gifts)", "dad")}
        if "air track" in QX[q]:
            sport = "gymnastics"
        if QX[q] in OVR:
            group, board = OVR[QX[q]]
        # finer tags from the title
        tl = t.lower()
        if not sport:
            for s in ["softball", "baseball", "basketball", "soccer", "football", "volleyball", "hockey", "lacrosse", "pickleball",
                      "tennis", "golf", "gymnastic", "tumbl", "cheer", "dance", "wrestling", "swim"]:
                if s in tl:
                    sport = {"gymnastic": "gymnastics", "tumbl": "gymnastics"}.get(s, s)
                    break
        if "golf" in tl and group not in ("Player gear (gift-worthy)",):
            group, board = ("Golf (sports dad gifts)", "dad")
        rows.append(dict(asin=a, name=t, price=p, rating=r, reviews=n, bought=b, query=QX[q], group=group,
                         sport=sport.title() if sport else "All sports", board=BOARDS[board],
                         tier=("$1,000+" if p >= 1000 else "$500-999" if p >= 500 else "$250-499" if p >= 250 else "$100-249"),
                         in_plan="Yes" if a in in_plan else ""))

if __name__ == "__main__":
    from collections import Counter
    print(len(rows), Counter(r["group"] for r in rows).most_common(), Counter(r["tier"] for r in rows), sum(r["in_plan"] == "Yes" for r in rows))


def build(out="/home/claude/Sideline_Ready_Big_Ticket_Products.xlsx"):
    F = "Arial"
    NAVY = "14213D"
    HF = Font(name=F, bold=True, color="FFFFFF", size=10)
    H = PatternFill("solid", start_color=NAVY)
    BF = Font(name=F, size=10)
    LF = Font(name=F, size=10, color="0563C1", underline="single")
    wb = Workbook()

    def table(ws, headers, data, widths, name, link_col=None):
        ws.append(headers)
        for r in data:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill = HF, H
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = BF
                if link_col and c.column == link_col:
                    c.hyperlink = c.value
                    c.font = LF
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        t = Table(displayName=name, ref=f"A1:{get_column_letter(len(headers))}{len(data)+1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "C2"
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            for c in row:
                c.number_format = '"$"#,##0.00'

    order = sorted(rows, key=lambda r: (r["group"], r["sport"], -num(r["bought"]), -r["reviews"]))
    hdr = ["#", "Product", "Price (Oct 30)", "Stars", "Reviews", "Bought last month", "Price tier", "Category", "Sport",
           "Pinterest board", "Found by searching", "Already in 365 plan", "ASIN", "Affiliate link"]
    W = [5, 58, 12, 7, 9, 11, 11, 26, 13, 38, 28, 11, 13, 50]
    def line(i, r):
        return [i, r["name"], r["price"], r["rating"], r["reviews"], r["bought"] or "-", r["tier"], r["group"], r["sport"],
                r["board"], r["query"], r["in_plan"], r["asin"], f"https://www.amazon.com/dp/{r['asin']}?tag={TAG}"]

    ws = wb.active
    ws.title = "Start Here"
    best = [r for r in order if num(r["bought"]) >= 100 and r["rating"] >= 4.5]
    lines = [
        ("Sideline Ready - Big-Ticket Amazon Products", 14, True),
        (f"{len(rows)} products found on Amazon on Oct 30, 2026 from 154 sports-mom searches (2 pages each, Amazon's $100+ price filter).", 10, False),
        ("", 10, False),
        ("Rules every product passed", 11, True),
        ("Price $100 or more  |  4.0 stars or higher  |  100+ customer reviews  |  not a sponsored ad", 10, False),
        ("Off-topic results were removed (home space heaters, men's work boots, bed blankets, etc.), and for very crowded searches", 10, False),
        ("(weighted vests, dumbbells, massage guns, air tracks) only the top 10-12 by sales and reviews were kept.", 10, False),
        ("", 10, False),
        ("Tabs", 11, True),
        ("All Products: every product, sorted by category, then sport, then sales.", 10, False),
        (f"Best Bets: the {len(best)} products that also show 100+ bought last month AND 4.5+ stars. Start your bundles here.", 10, False),
        ("Summary: live counts by category, sport and price tier.", 10, False),
        ("", 10, False),
        ("Good to know", 11, True),
        ("'Bought last month' is Amazon's own badge. A dash means Amazon didn't show one (usually under 50 sales that month).", 10, False),
        ("Prices change daily. Check the price before you write it into a post (and never hard-code prices on the site).", 10, False),
        ("'Already in 365 plan' = this product is already one of the 365 daily anchors.", 10, False),
        ("Links use your tracking ID screenfreet0f-20.", 10, False),
    ]
    for i, (t, size, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=F, size=size, bold=bold, color=NAVY if bold else "000000")
    ws.column_dimensions["A"].width = 125

    ws = wb.create_sheet("All Products")
    table(ws, hdr, [line(i, r) for i, r in enumerate(order, 1)], W, "AllProducts", link_col=14)
    ws = wb.create_sheet("Best Bets")
    best = sorted(best, key=lambda r: (-num(r["bought"]), -r["reviews"]))
    table(ws, hdr, [line(i, r) for i, r in enumerate(best, 1)], W, "BestBets", link_col=14)

    ws = wb.create_sheet("Summary")
    ws["A1"] = "Summary"
    ws["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
    r0 = 3
    for title, col, vals in [("By category", "H", sorted({r["group"] for r in rows})),
                             ("By sport", "I", sorted({r["sport"] for r in rows})),
                             ("By price tier", "G", ["$100-249", "$250-499", "$500-999", "$1,000+"])]:
        ws.cell(row=r0, column=1, value=title).font = Font(name=F, bold=True, color=NAVY)
        for c, v in zip((1, 2, 3), ("Group", "Products", "Best Bets")):
            x = ws.cell(row=r0 + 1, column=c, value=v)
            x.font, x.fill = HF, H
        start = r0 + 2
        for i, v in enumerate(vals, start):
            ws.cell(row=i, column=1, value=v).font = BF
            ws.cell(row=i, column=2, value=f"=COUNTIF('All Products'!{col}:{col},A{i})").font = BF
            ws.cell(row=i, column=3, value=f"=COUNTIF('Best Bets'!{col}:{col},A{i})").font = BF
        end = start + len(vals)
        ws.cell(row=end, column=1, value="Total").font = Font(name=F, bold=True)
        ws.cell(row=end, column=2, value=f"=SUM(B{start}:B{end-1})").font = Font(name=F, bold=True)
        ws.cell(row=end, column=3, value=f"=SUM(C{start}:C{end-1})").font = Font(name=F, bold=True)
        r0 = end + 2
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    wb.save(out)
    return out, len(order), len(best)


if __name__ == "__main__":
    print(build())
