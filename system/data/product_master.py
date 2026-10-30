import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
sys.path.insert(0, "/home/claude")
from data_posts import POSTS
from selected import SEL
from sl_data import P as SLP, BOARDS as SLB, board_for
from sl_data2 import P2

TAG = "screenfreet0f-20"
link = lambda a: f"https://www.amazon.com/dp/{a}?tag={TAG}"
SEASON = {1: "Nov-Feb (holiday & winter)", 2: "Year-round", 3: "Mar-Aug (spring & summer)"}

# ---------- Dr. Q Kids ----------
pins = json.load(open("/home/claude/pins150.json"))
boards = {b["key"]: b["name"] for b in json.load(open("/home/claude/assets_build/boards.json"))}
TIER1 = {1,2,3,4,5,6,7,8,12,13,14,22,30,31,32,33,46,48,50,51,52}
TIER3 = {20,23,24,25,26,28,29,35,38,45,47}
tier_of = lambda post: 1 if post in TIER1 else 3 if post in TIER3 else 2
first_kw = {}
for p in pins:
    first_kw.setdefault(p["asin"], (p["kw"], p["board"]))
kids = []
for post, items in SEL.items():
    P_ = POSTS[post - 1]
    for a, name, price, rating, nr, bought in items:
        if a in {k[0] for k in kids}: continue
        kw, board = first_kw.get(a, (P_[1], boards.get(P_[3], P_[3])))
        kids.append([a, name, f"${price:.2f}" if isinstance(price, (int, float)) else price, rating, nr, bought,
                     tier_of(post), kw, P_[1], "; ".join(P_[2]), board])
# 6 new (Oct 30 scan)
NEW6 = [
 ("B00008BFZH","Snap Circuits Explorer SC-100 STEM Kit","$32.99",4.8,"29.4K","10K+",1,"stem gifts for kids","screen free christmas gifts","stem toys for 8 year olds; screen free gifts for kids; educational gifts for kids","gifts"),
 ("B07R92KFZR","National Geographic Glowing Marble Run","$37.99",4.7,"14.2K","3K+",1,"marble run for kids","screen free activities for kids indoor","indoor activities for kids; screen free gifts for kids; stem toys for 5 year olds","screenfree"),
 ("B078Y44VTS","Melissa & Doug Water Wow! Dinosaurs","$8.23",4.8,"12.4K","4K+",2,"road trip activities for toddlers","road trip activities for kids 3-5","quiet activities for toddlers; mess free activities for toddlers; restaurant activities for kids","travel"),
 ("B098PMWSVS","Cupkin Animal Sticker Book + Coloring Pages","$16.98",4.8,"7K","7K+",2,"quiet time activities for kids","quiet time boxes","sticker books for kids; quiet time box ideas; screen free activities for kids","calm"),
 ("B00TPMDNSU","Play-Doh Starter Set (9 pc)","$7.12",4.8,"10.3K","20K+",2,"rainy day activities for kids","rainy day activities for kids","play doh activities; indoor activities for toddlers; stocking stuffers for kids 3-5","screenfree"),
 ("B00272N8L2","Melissa & Doug Primary Lacing Beads (33 pc)","$19.99",4.8,"21K","2K+",2,"montessori toys for 2 year olds","montessori toys 2-3","fine motor activities for toddlers; quiet activities for toddlers; busy toys for toddlers","montessori"),
]
for r in NEW6:
    kids.append(list(r[:10]) + [boards.get(r[10], r[10])])
assert len(kids) == 150 and len({k[0] for k in kids}) == 150, len(kids)

# ---------- Sideline Ready ----------
slb = {k: n for k, n, _ in SLB}
side = []
for a, name, price, rating, nr, bought, tier, kws in SLP:
    side.append([a, name, f"${price:.2f}" if isinstance(price, (int, float)) else price, rating, nr, bought, tier, kws[0], "; ".join(kws[1:]), slb[board_for(kws[0])], "Nov 2026 pins"])
for a, name, price, rating, nr, bought, tier, kws, bk in P2:
    side.append([a, name, price, float(rating), nr, bought, tier, kws[0], "; ".join(kws[1:]), slb[bk], "New - from Dec 2026"])
assert len(side) == 150 and len({s[0] for s in side}) == 150, len(side)

# ---------- TPT ----------
cat = json.load(open("/home/claude/quenwade44/drq-pins/system/catalog.json"))
tb = cat["boards"]
TYPE = {"doc": "Bundle / kit / library", "mine": "Original tool / form", "test": "Test report template"}
PRI = {"doc": 1, "mine": 2, "test": 3}
tpt = sorted(cat["products"], key=lambda p: (PRI[p["kind"]], p["board"], p["name"]))
groupA = {p["id"] for p in tpt[:150]}
repeat = {p["id"] for p in tpt[:80]}  # bundles + top originals also run in month B

# ---------- workbook ----------
F = "Arial"
H = PatternFill("solid", start_color="1E3A5F")
HF = Font(name=F, bold=True, color="FFFFFF", size=10)
BF = Font(name=F, size=10)
LF = Font(name=F, size=10, color="0563C1", underline="single")
thin = Side(style="thin", color="D0D7E2")
wb = Workbook()

def sheet(ws, headers, rows, widths, link_cols=(), tname="T"):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = HF, H
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF
            c.alignment = Alignment(vertical="top", wrap_text=False)
            if c.column in link_cols and c.value:
                c.hyperlink = c.value
                c.font = LF
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    t = Table(displayName=tname, ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)

# README
ws = wb.active; ws.title = "Start Here"
lines = [
 ("Product Master List - all 3 Pinterest accounts", True),
 ("Built Oct 30, 2026 for Dr. Q. One row = one product. Each month = 150 pins per account (5 a day).", False),
 ("", False),
 ("Tab", True),
 ("TPT - School Templates: every product publicly listed in the TPT store, the Pinterest board it pins to, and which monthly rotation it runs in.", False),
 ("Amazon - Dr Q Kids: 150 Amazon products for Dr. Q | Screen-Free Kid Toys, with the exact parent search phrase for each pin and the affiliate link.", False),
 ("Amazon - Sideline Ready: 150 Amazon products for Sideline Ready (50 already pinned in November + 100 new).", False),
 ("Summary: live counts by board, season and rotation (formulas - they update if you edit the tabs).", False),
 ("", False),
 ("How the Amazon products were picked", True),
 ("Every Amazon product showed '1K+ bought in past month' or more on Amazon search, was not a sponsored listing, and rated 4.2 stars or higher when checked (Oct 2026).", False),
 ("Each product is matched to a phrase parents actually type into Pinterest search - that phrase goes in the pin title, description and alt text.", False),
 ("Season column = when to push it hardest. Year-round products can run every month; winter products Nov-Feb; summer products Mar-Aug.", False),
 ("Amazon sales counts and prices change. Each month before building pins I re-check that every product still has 1K+ sales and swap any that dropped.", False),
 ("", False),
 ("How TPT rotation works (150 pins a month from 220 products)", True),
 ("Rotation A (Dec, Feb, Apr...): 150 products - all bundles/kits, all original tools, and the first test report templates.", False),
 ("Rotation B (Jan, Mar, May...): the other 70 test report templates + the top 80 (bundles and originals) again with fresh pin designs.", False),
 ("Result: bundles get a pin every month, and every product gets at least one pin every 2 months.", False),
 ("", False),
 ("Things to check", True),
 ("TPT: only 220 products are publicly listed on the store page - you mentioned 236. The other 16 may be drafts, unlisted, or very new. Send me their names or links and I'll add them.", False),
 ("Sideline Ready links use your current tracking ID (screenfreet0f-20). If you make a separate ID for Sideline Ready, send it and I'll swap every link.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=14 if i == 1 else 10, bold=b, color="1E3A5F" if b else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120

# TPT
ws = wb.create_sheet("TPT - School Templates")
rows = []
for i, p in enumerate(tpt, 1):
    name = p["name"] + (f" | {p['sub']}" if p.get("sub") else "")
    rot = "A + B (every month)" if p["id"] in repeat else ("A" if p["id"] in groupA else "B")
    rows.append([i, name, p["fmt"], TYPE[p["kind"]], tb[str(p["board"])], rot, p["id"], p["url"]])
sheet(ws, ["#", "Product", "Format", "Type", "Pinterest board", "Rotation", "TPT ID", "TPT link"],
      rows, [5, 70, 18, 22, 48, 20, 11, 60], link_cols=(8,), tname="TPT")

# Kids
ws = wb.create_sheet("Amazon - Dr Q Kids")
rows = [[i, k[1], k[2], k[3], k[4], k[5], SEASON[k[6]], k[7], k[8], k[9], k[10], k[0], link(k[0])] for i, k in enumerate(kids, 1)]
sheet(ws, ["#", "Product", "Price (Oct 30)", "Stars", "Ratings", "Bought last month", "Season", "Pin search phrase",
           "Topic", "More parent search phrases", "Pinterest board", "ASIN", "Affiliate link"],
      rows, [5, 46, 12, 7, 9, 12, 24, 34, 30, 60, 36, 13, 52], link_cols=(13,), tname="Kids")

# Sideline
ws = wb.create_sheet("Amazon - Sideline Ready")
rows = [[i, s[1], s[2], s[3], s[4], s[5], SEASON[s[6]], s[7], s[8], s[9], s[10], s[0], link(s[0])] for i, s in enumerate(side, 1)]
sheet(ws, ["#", "Product", "Price (Oct 30)", "Stars", "Ratings", "Bought last month", "Season", "Pin search phrase",
           "More parent search phrases", "Pinterest board", "Status", "ASIN", "Affiliate link"],
      rows, [5, 46, 12, 7, 9, 12, 24, 34, 60, 38, 20, 13, 52], link_cols=(13,), tname="Sideline")

# Summary with formulas
ws = wb.create_sheet("Summary")
r = 1
def head(t):
    global r
    c = ws.cell(row=r, column=1, value=t); c.font = Font(name=F, bold=True, size=12, color="1E3A5F"); r += 1
def block(title, sheetname, col, values, total_label="Total"):
    global r
    head(title)
    for c, v in zip((1, 2), ("Group", "Products")):
        x = ws.cell(row=r, column=c, value=v); x.font, x.fill = HF, H
    r += 1; start = r
    for v in values:
        ws.cell(row=r, column=1, value=v).font = BF
        ws.cell(row=r, column=2, value=f"=COUNTIF('{sheetname}'!${col}:${col},A{r})").font = BF
        r += 1
    ws.cell(row=r, column=1, value=total_label).font = Font(name=F, bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B{start}:B{r-1})").font = Font(name=F, bold=True)
    r += 2
block("TPT - products by rotation", "TPT - School Templates", "F", ["A", "B", "A + B (every month)"])
block("TPT - products by board", "TPT - School Templates", "E", [tb[str(i)] for i in range(1, 13)])
block("Dr Q Kids - products by season", "Amazon - Dr Q Kids", "G", list(SEASON.values()))
block("Dr Q Kids - products by board", "Amazon - Dr Q Kids", "K", sorted({k[10] for k in kids}))
block("Sideline Ready - products by season", "Amazon - Sideline Ready", "G", list(SEASON.values()))
block("Sideline Ready - products by board", "Amazon - Sideline Ready", "J", [n for _, n, _ in SLB])
block("Sideline Ready - status", "Amazon - Sideline Ready", "K", ["Nov 2026 pins", "New - from Dec 2026"])
ws.column_dimensions["A"].width = 58; ws.column_dimensions["B"].width = 12

out = "/home/claude/Product_Master_List_3_Channels.xlsx"
wb.save(out); print(out, len(tpt), len(kids), len(side))
