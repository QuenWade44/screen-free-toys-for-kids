"""Day 1 (Sat Oct 31, 2026): 15 Sideline Ready pins linking straight to Amazon."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TAG = "screenfreet0f-20"
# asin, product, short, price, stars, reviews, bought, phrase, phrase2, phrase3, board, what-it-does, scene
P = [
("B07H7SN831","Pocket Radar Smart Coach","Pocket Radar Smart Coach",399.00,4.6,"1,400","500+","baseball dad gifts","softball dad gifts","baseball coach christmas gifts","Sports Dad Essentials & Gifts",
 "a pocket-size radar that reads pitch speed and exit velo, with an app that saves every rep","dad behind a backstop at dusk holding a small black device, teen pitcher blurred in background"),
("B0CNS894RH","Theragun Relief Massage Gun","Theragun Relief",159.99,4.5,"2,600","6K+","christmas gifts for teen athletes","christmas gifts athletes","gifts for athletes","Sports Mom Gifts & Gift Baskets",
 "a lightweight massage gun for sore legs after practice and long tournament weekends","wrapped gift box under a Christmas tree next to a teen's cleats and gym bag"),
("B0FS8T1WQW","Lifetime NBA League Edition Portable Basketball Hoop","Lifetime portable hoop",269.00,4.6,"5,400","100+","backyard basketball hoop","backyard basketball court ideas","basketball christmas gifts ideas","Backyard Sports Training Equipment",
 "an adjustable-height portable hoop that grows with your player from driveway games to real reps","suburban driveway at golden hour, portable basketball hoop, kid shooting a jump shot, no logos"),
("B0GR9YG1LY","DJI Osmo Mobile 8 Gimbal + Mic Mini","DJI Osmo Mobile 8",182.00,4.6,"14,200","1K+","sports filming","sports camera shots","christmas gifts for sports mom","Sports Mom Gifts & Gift Baskets",
 "a phone gimbal that tracks your kid down the field and keeps the video smooth, plus a clip-on mic","mom in the bleachers filming a youth soccer game on a phone mounted on a small gimbal"),
("B0CT3SGHXL","Garmin Forerunner 165 Running Smartwatch","Garmin Forerunner 165",192.49,4.6,"5,300","7K+","christmas gifts for track runners","gifts for girls in track","track gifts ideas","Sports Mom Gifts & Gift Baskets",
 "a GPS running watch that tracks pace, distance and training load for cross country and track kids","close-up of a runner's wrist at a high school track at sunrise, generic black sports watch"),
("B0C2YCMB3R","HOPERAN 25-Inch Dual-Control Heated Stadium Seat","heated stadium seat",129.99,4.5,"679","400+","stadium seats for bleachers","cold football games","sports mom cold weather essentials","Stadium Seats & Bleacher Comfort",
 "a padded bleacher seat with back support and two heat zones, powered by a battery pack","navy padded stadium seat on metal bleachers under Friday night lights, breath visible in cold air"),
("B01H50RF46","ORORO Women's Heated Vest","ORORO heated vest",129.99,4.5,"11,100","1K+","sports mom cold weather essentials","cold football game fits","winter sports mom essentials","Cold Football Game Essentials",
 "a lightweight battery-heated vest that layers under a jacket for cold games and early tournaments","mom in a quilted vest and beanie on bleachers at a night football game, warm stadium lights"),
("B082HHTKB6","UGG Women's Classic Ultra Mini Boot","UGG Ultra Mini",164.95,4.5,"10,400","1K+","christmas gifts for sports mom","sports mom christmas gift","gifts for the sports mom","Sports Mom Gifts & Gift Baskets",
 "the cozy ankle-height boot that keeps her feet warm on cold bleachers and slips on in seconds","chestnut suede ankle boots on a cozy blanket next to a stadium cup and team scarf, gift-wrap ribbon"),
("B07ZS5RQZ4","Blackstone 22-Inch Tabletop Griddle with Hood","Blackstone 22\" griddle",219.00,4.7,"2,800","1K+","tailgate setup ideas","tailgate food ideas","sports dad gifts","Sports Dad Essentials & Gifts",
 "a portable propane griddle with a hood that cooks breakfast or burgers for the whole team at a tailgate","tailgate tent at a youth football field, tabletop griddle sizzling pancakes and bacon, families in team colors"),
("B082TMBYR6","Jackery Explorer 300 Portable Power Station","Jackery Explorer 300",289.00,4.6,"11,300","8K+","travel ball must haves","travel ball essentials","tournament day essentials","Travel Ball & Tournament Must Haves",
 "a quiet battery power station that charges phones, fans, heated seats and a tablet all tournament long","tournament canopy at a baseball complex, compact orange power station charging phones and a fan, wagon nearby"),
("B07Q82MG8S","Mr. Heater 18,000 BTU Big Buddy Heater","Mr. Heater Big Buddy",155.98,4.6,"5,700","900+","cold football games","tailgate ideas","winter sports mom essentials","Cold Football Game Essentials",
 "a portable propane heater for tailgates and open-air sidelines (check your field's rules first)","family tailgate at dusk in winter coats around a red portable heater, steam from hot cocoa cups"),
("B0G23XKMJL","Timber Ridge 400L Folding Double Decker Wagon","double decker sports wagon",124.99,4.6,"3,600","1K+","sports wagon ideas","sports wagon organization","sports mom wagon","Sports Wagon Ideas & Accessories",
 "a big folding wagon with a second shelf and brakes, so chairs, coolers and gear make it to the field in one trip","loaded double decker wagon rolling across a soccer complex parking lot, chairs and cooler on top, morning light"),
("B078815CXR","FBSPORT Inflatable Air Track Tumbling Mat","air track tumbling mat",119.99,4.7,"3,100","400+","gymnastics gifts for kids","tumbling mats for kids","gymnastics equipment for home","Backyard Sports Training Equipment",
 "an inflatable tumbling track for cartwheels and back walkovers at home, rolls up when practice is over","living room with a long inflatable tumbling mat, girl mid-cartwheel, holiday decorations in background"),
("B07XFPSZ58","Pop-A-Shot Dual Shot Sport Arcade Basketball","Pop-A-Shot Dual Shot",299.99,4.6,"3,300","200+","basketball christmas gifts ideas","basketball gifts for boys","sports room for boys","Sports Mom Gifts & Gift Baskets",
 "a two-hoop electronic arcade basketball game for the basement or garage, built for nonstop shooting","basement game room, two-hoop arcade basketball machine glowing, two siblings shooting, Christmas lights"),
("B0FRNG2N5H","BowFlex SelectTech Adjustable Dumbbells","BowFlex SelectTech dumbbells",399.99,4.8,"334","3K+","home gym for teen boy","teen home gym ideas","garage gym ideas","Backyard Sports Training Equipment",
 "adjustable dumbbells that replace a full rack, so a teen athlete can lift at home in a small corner","clean garage gym corner with a bench and adjustable dumbbells, teen athlete doing curls, morning light"),
]
QH = [
("What do you get a baseball dad who already has everything?", "The gift he'll use every bullpen: a radar that clocks every pitch."),
("What Christmas gift will a teen athlete actually use?", "The recovery tool they'll grab after every practice."),
("Which backyard hoop grows with my kid?", "The height adjusts as they grow, so you only buy one hoop."),
("How do I film my kid's games without shaky video?", "Clip in your phone and it follows the play for you."),
("What do you get a runner who has everything?", "A GPS watch that tracks every mile, split and pace."),
("How do I survive 3 hours on metal bleachers?", "Back support plus two heat zones. Your back will thank you."),
("How do I stay warm at night games without 5 bulky layers?", "One thin battery-heated vest under your jacket. That's it."),
("What do I get the mom who lives in the bleachers?", "Cozy boots she can slip on before every cold game."),
("How do I feed the whole team at a tailgate?", "One portable griddle: pancakes at 7, burgers at noon."),
("How do I keep phones and fans charged all tournament weekend?", "A quiet power station that charges everything from your canopy."),
("How do we stay warm tailgating before a freezing game?", "A portable propane heater for the tailgate (check field rules first)."),
("How do I get chairs, cooler and gear to the field in one trip?", "A double-decker wagon: chairs on top, cooler below, one trip."),
("Where can my gymnast practice tumbling at home?", "An air track for cartwheels in the living room. Rolls up after."),
("What basketball gift keeps them off screens all winter?", "Two hoops, a scoreboard, and nonstop shooting in the basement."),
("How do I set up a teen home gym in a small space?", "One pair of adjustable dumbbells replaces a whole rack."),
]
TIMES = ["7:00 AM", "8:00 AM", "9:00 AM", "10:00 AM", "11:00 AM", "12:00 PM", "1:00 PM", "2:00 PM", "3:00 PM",
         "4:00 PM", "5:00 PM", "6:00 PM", "7:00 PM", "8:00 PM", "9:00 PM"]

def tc(s):
    small = {"for", "and", "of", "to", "a", "the", "in", "on"}
    return " ".join(w if (i and w in small) else w[:1].upper() + w[1:] for i, w in enumerate(s.split()))

rows = []
for i, (a, name, short, price, stars, rev, bought, k1, k2, k3, board, what, scene) in enumerate(P, 1):
    title = f"{tc(k1)}: {short}"
    desc = (f"{k1[:1].upper()+k1[1:]}: the {short} is {what}. "
            f"A smart pick if you're also searching {k2} or {k3}. "
            f"Tap to see it on Amazon. Affiliate link: I may earn a small commission at no extra cost to you.")
    alt = f"{scene[:1].upper()+scene[1:]}, with the words \"{k1}\" across the top."
    overlay = k1.upper()
    img = (f"Vertical 2:3 photo, {scene}. Natural light, candid, realistic. No brand logos, no readable text. "
           f"Leave the top third clear for a text overlay.")
    vid = (f"9:16 or 2:3, 8-10 seconds. Shot 1 (3s): your brand avatar at the field/bleachers, looks at camera and points up "
           f"to the overlay \"{k1}\". Shot 2 (5s): {scene}. Slow push-in. End card: product name + \"Tap the pin\".")
    line = (f"Searching {k1}? Here's one parents are buying: the {short}. Tap the pin to see it.")
    q, hook = QH[i - 1]
    rows.append([i, "Sat Oct 31", TIMES[i - 1], name, price, stars, rev, bought, overlay, q, hook, title, desc, alt, board,
                 f"https://www.amazon.com/dp/{a}?tag={TAG}", img, vid, line, f"day001-{i:02d}-{k1.replace(' ', '-')}", a])
    assert len(title) <= 100 and len(desc) <= 500, (title, len(desc))
    assert len(hook) <= 75, hook

wb = Workbook()
ws = wb.active
ws.title = "Day 1 - 15 Pins"
hdr = ["#", "Day", "Post time (CT)", "Product", "Price (Oct 30)", "Stars", "Reviews", "Bought last month",
       "Text overlay (exact search phrase)", "The question she's asking", "Hook (answers it, on the pin under the overlay)", "Pin title", "Pin description", "Alt text", "Pinterest board", "Amazon affiliate link",
       "Image prompt (OpenArt / Runway)", "Video plan (8-10 sec)", "Avatar line (caption or voice)", "File name", "ASIN"]
ws.append(hdr)
for r in rows:
    ws.append(r)
F = "Arial"
for c in ws[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", start_color="14213D")
    c.alignment = Alignment(wrap_text=True, vertical="center")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(vertical="top", wrap_text=c.column in (10, 11, 12, 13, 14, 17, 18, 19))
        if c.column == 16:
            c.hyperlink = c.value
            c.font = Font(name=F, size=10, color="0563C1", underline="single")
        if c.column == 5:
            c.number_format = '"$"#,##0.00'
W = [4, 10, 10, 34, 10, 6, 9, 10, 30, 38, 42, 40, 60, 45, 30, 42, 55, 55, 45, 34, 13]
for i, w in enumerate(W, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
t = Table(displayName="Day1", ref=f"A1:{get_column_letter(len(hdr))}{len(rows)+1}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
ws.add_table(t)
ws.freeze_panes = "E2"
ws2 = wb.create_sheet("Rules")
rules = [
 "Every product: $100+, 4.0+ stars, 100+ reviews, not an ad. Today's 15 also sold 100+ last month and have 4.5+ stars (Best Bets).",
 "Text overlay = the exact search phrase, word for word. Pin title starts with it. Description uses it first, then 2 related phrases.",
 "Every pin names the question she has when she types that phrase, and the hook (one line under the overlay) answers it. Keep hooks under 75 characters.",
 "No prices, star ratings or review counts on the pin or in the voiceover (Amazon only allows those when pulled live from Amazon).",
 "No Amazon product photos. Use your own photos, brand-approved media, or AI lifestyle scenes like the prompts here.",
 "Your avatar is a host, not a customer: it can point, explain and say 'parents are buying'. It never says 'I bought/use/love this'.",
 "Add a small 'AI-generated' note on any pin that uses your AI avatar.",
 "Link straight to Amazon with your tag (screenfreet0f-20). Prices were checked Oct 30; re-check before you post.",
]
for i, r in enumerate(rules, 1):
    c = ws2.cell(row=i, column=1, value=f"{i}. {r}")
    c.font = Font(name=F, size=10)
ws2.column_dimensions["A"].width = 140
wb.save("/home/claude/Sideline_Ready_Day1_15_Pins.xlsx")
print("ok", len(rows), max(len(r[11]) for r in rows), max(len(r[12]) for r in rows), max(len(r[10]) for r in rows))
